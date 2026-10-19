"""Fonctions techniques du test USB : détection, E/S, erreurs et rapports.
Ce module ne gère ni les threads de test ni l'interface graphique.
"""

import os
import sys
import json
import errno
import mmap
import shutil
import fcntl
import struct
import tempfile
import subprocess
import numpy as np
import pyudev
from io import BytesIO
from pathlib import Path
from docx import Document
from docx.shared import Inches
from datetime import datetime
from config import (
    USB_TEST_PORT, RESULTS_DIR, MIN_WRITE_MBPS, 
    MIN_READ_MBPS, MAX_READ_MBPS, TEST_LINK_FILE, 
    MAX_MATING_CYCLES
    )


BLKGETDISKSEQ = 0x80081280

# ============================================================
# 1. OUTILS SYSTÈME
# ============================================================

def run_command(cmd, timeout=20, input_text=None):
    """Exécute une commande et retourne (stdout, stderr, returncode)."""
    try:
        result = subprocess.run(
            cmd,
            input=input_text,
            capture_output=True,
            text=True,
            timeout=timeout
        )
        return result.stdout.strip(), result.stderr.strip(), result.returncode

    except (OSError, subprocess.TimeoutExpired) as e:
        return "", str(e), -1

    
def get_root_disk():
    """Retourne le nom du disque contenant la partition système."""
    source, _, rc = run_command(["findmnt", "-n", "-o", "SOURCE", "/"])

    if rc != 0 or not source:
        return None

    disk, _, rc = run_command(["lsblk", "-no", "PKNAME", source])

    return disk.strip() if rc == 0 and disk else None


ROOT_DISK = get_root_disk()

# ============================================================
# 2. DÉTECTION ET IDENTIFICATION USB
# ============================================================

def scan_usb_candidates():
    """Liste les disques USB connectés au port de test, hors disque système."""
    candidates = []

    for dev in pyudev.Context().list_devices(subsystem="block", DEVTYPE="disk"):
        if dev.get("ID_BUS") != "usb":
            continue

        if not dev.get("ID_PATH", "").startswith(USB_TEST_PORT + ":"):
            continue

        if dev.sys_name == ROOT_DISK:
            continue

        candidates.append(dev.device_node)

    return sorted(candidates)


def get_usb_parent_device(udev_device):
    """Retourne le périphérique USB parent (VID/PID)."""
    return udev_device.find_parent("usb", "usb_device")


def get_diskseq(devnode):
    """Retourne le DISKSEQ du périphérique pour identifier son instance."""
    try:
        device = pyudev.Devices.from_device_file(pyudev.Context(), devnode)
        return int(device.properties["DISKSEQ"])

    except (KeyError, ValueError, OSError, pyudev.DeviceNotFoundError) as exc:
        raise RuntimeError(
            f"Unable to identify the device {devnode}"
        ) from exc


def assert_same_device(devnode, expected_diskseq):
    """Vérifie que le périphérique n'a pas été remplacé."""
    current_diskseq = get_diskseq(devnode)

    if current_diskseq != expected_diskseq:
        raise RuntimeError(
            f"Device replaced on {devnode} : "
            f"DiskSeq expected={expected_diskseq}, actual={current_diskseq}"
        )


def assert_test_port(devnode):
    """Vérifie que le périphérique utilise le port USB réservé."""
    dev = pyudev.Devices.from_device_file(pyudev.Context(), devnode)
    id_path = dev.get("ID_PATH", "")

    if not id_path.startswith(USB_TEST_PORT + ":"):
        raise RuntimeError(
            f"Unauthorized port : {devnode} (ID_PATH={id_path})"
        )

# ============================================================
# 3. INFORMATIONS DU PÉRIPHÉRIQUE
# ============================================================

def get_device_info(devnode):
    """Récupère les informations du périphérique via lsblk et pyudev."""
    info = {
        "Devnode": devnode,
        "Model": "N/A",
        "Manufacturer": "N/A",
        "Serial": "N/A",
        "Size_GB": "N/A",
        "VID": "N/A",
        "PID": "N/A",
        "USBVersion": "N/A",
        "Date": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }

    # Informations du disque via lsblk
    lb = get_lsblk_info(devnode)

    info["Model"] = lb.get("model") or "N/A"
    info["Manufacturer"] = lb.get("vendor") or "N/A"
    info["Serial"] = lb.get("serial") or "N/A"

    try:
        if lb.get("size"):
            info["Size_GB"] = round(int(lb["size"]) / 1e9, 2)
    except (ValueError, TypeError):
        pass

    # Informations USB via pyudev
    try:
        udev_dev = pyudev.Devices.from_device_file(pyudev.Context(), devnode)
        usb_dev = get_usb_parent_device(udev_dev)

        if usb_dev is None:
            return info

        info["VID"] = (
            (usb_dev.attributes.get("idVendor") or b"")
            .decode(errors="ignore").upper() or "N/A"
        )
        info["PID"] = (
            (usb_dev.attributes.get("idProduct") or b"")
            .decode(errors="ignore").upper() or "N/A"
        )

        if info["Manufacturer"] == "N/A":
            info["Manufacturer"] = (
                (usb_dev.attributes.get("manufacturer") or b"")
                .decode(errors="ignore").strip() or "N/A"
            )

        if info["Model"] == "N/A":
            info["Model"] = (
                (usb_dev.attributes.get("product") or b"")
                .decode(errors="ignore").strip() or "N/A"
            )

        speed = (
            (usb_dev.attributes.get("speed") or b"")
            .decode(errors="ignore").strip()
        )
        info["USBVersion"] = get_usb_speed_label(speed) or "N/A"

    except (OSError, pyudev.DeviceNotFoundError, AttributeError, UnicodeError):
        pass

    return info


def get_lsblk_info(devnode):
    """Retourne les informations du disque fournies par lsblk."""
    out, _, rc = run_command([
        "lsblk", "-b", "-J", "-o", "NAME,MODEL,SERIAL,SIZE,VENDOR", devnode
    ])

    if rc != 0 or not out:
        return {}

    try:
        return json.loads(out)["blockdevices"][0]
    except (json.JSONDecodeError, KeyError, IndexError, TypeError):
        return {}


def get_usb_speed_label(speed_mbps):
    """Convertit une vitesse USB en Mb/s en libellé lisible."""
    try:
        speed = float(speed_mbps)
    except (TypeError, ValueError):
        return "N/A"

    labels = {
        1.5: "USB Low-Speed (1.5 Mb/s)",
        12: "USB Full-Speed (12 Mb/s)",
        480: "USB 2.0 High-Speed (480 Mb/s)",
        5000: "USB 3.2 Gen 1 (5 Gb/s)",
        10000: "USB 3.2 Gen 2 (10 Gb/s)",
        20000: "USB 3.2 Gen 2x2 (20 Gb/s)",
    }

    for reference, label in labels.items():
        if abs(speed - reference) < reference * 0.05:
            return label

    return f"{speed:.0f} Mb/s"


def get_device_size_bytes(devnode):
    """Retourne la capacité du périphérique bloc en octets."""
    out, _, rc = run_command(["blockdev", "--getsize64", devnode])

    if rc == 0 and out.isdigit():
        return int(out)

    fd = os.open(devnode, os.O_RDONLY)
    try:
        return os.lseek(fd, 0, os.SEEK_END)
    finally:
        os.close(fd)

# ============================================================
# 4. ACCÈS DIRECT AU PÉRIPHÉRIQUE
# ============================================================

def create_aligned_buffer(size, fill=None):
    """Crée un buffer mmap aligné en mémoire, compatible avec O_DIRECT."""
    buf = mmap.mmap(-1, size)

    if fill is not None:
        buf[:] = fill

    return buf


def open_test_device(devnode, expected_diskseq):
    """Ouvre le disque en E/S directes et vérifie son identité."""
    assert_test_port(devnode)
    assert_same_device(devnode, expected_diskseq)

    fd = os.open(
        devnode,
        os.O_RDWR | os.O_SYNC | os.O_DIRECT | os.O_EXCL
    )

    try:
        diskseq = struct.unpack(
            "Q", fcntl.ioctl(fd, BLKGETDISKSEQ, bytes(8))
        )[0]

        if diskseq != expected_diskseq:
            raise RuntimeError(
                "The USB device was changed before it was opened"
            )

        return fd

    except Exception:
        os.close(fd)
        raise

# ============================================================
# 5. GESTION DES PARTITIONS ET REFORMATAGE
# ============================================================

FILESYSTEM_CONFIG = {
        "exfat": {
            "tool": "mkfs.exfat",
            "args": [],
            "label_arg": "-n",
            "blkid": "exfat",
            "package": "exfatprogs",
        },
        "fat32": {
            "tool": "mkfs.vfat",
            "args": ["-F", "32"],
            "label_arg": "-n",
            "blkid": "vfat",
            "package": "dosfstools",
        },
        "ntfs": {
            "tool": "mkfs.ntfs",
            "args": ["-Q"],
            "label_arg": "-n",
            "blkid": "ntfs",
            "package": "ntfs-3g",
        },
    }


def get_mounted_devices(devnode, expected_diskseq):
    """Retourne les périphériques montés appartenant au disque USB."""
    assert_same_device(devnode, expected_diskseq)

    out, err, rc = run_command([
        "lsblk", "-J", "-p", "-l",
        "-o", "PATH,TYPE,MOUNTPOINT", devnode
    ])

    if rc != 0:
        raise RuntimeError(f"Unable to read partitions : {err or out}")

    return [
        device["path"]
        for device in json.loads(out)["blockdevices"]
        if device.get("mountpoint")
    ]


def unmount_device(devnode, expected_diskseq):
    """Démonte tous les systèmes de fichiers montés sur le disque USB."""
    for device in reversed(get_mounted_devices(devnode, expected_diskseq)):
        assert_same_device(devnode, expected_diskseq)

        out, err, rc = run_command(["umount", device])
        if rc != 0:
            raise RuntimeError(
                f"Unable to unmount{device} : {err or out}"
            )

    remaining = get_mounted_devices(devnode, expected_diskseq)
    if remaining:
        raise RuntimeError(
            f"Devices still mounted ; {', '.join(remaining)}"
        )


def execute_device_command(devnode, expected_diskseq, cmd, timeout=30, input_text=None):
    """Exécute une commande après vérification du disque et du port USB."""
    assert_same_device(devnode, expected_diskseq)
    assert_test_port(devnode)

    out, err, rc = run_command(cmd, timeout=timeout, input_text=input_text)

    if rc != 0:
        raise RuntimeError(
            f"Command Failed : {' '.join(cmd)}\n{err or out}"
        )

    return out


def reformat_device(devnode, fs_type, label="", expected_diskseq=None):
    """Recrée une table MBR, une partition unique et son système de fichiers."""

    if expected_diskseq is None:
        raise RuntimeError("Reformatting prohibited: DiskSeq not provided")

    fs_type = fs_type.lower().strip()

    if fs_type not in FILESYSTEM_CONFIG:
        raise ValueError(f"Unsupported file system : {fs_type}")

    cfg = FILESYSTEM_CONFIG[fs_type]

    # Vérifier les dépendances
    for tool in (cfg["tool"], "sfdisk", "partprobe", "udevadm", "lsblk", "blkid", "sync"):
        if not shutil.which(tool):
            raise RuntimeError(f"Tool not found : {tool}")

    # Créer une table MBR avec une partition unique
    partition_type = {
        "fat32": "c",
        "exfat": "7",
        "ntfs": "7",
    }[fs_type]

    script = f"label: dos\nstart=2048, type={partition_type}\n"

    execute_device_command(
        devnode, expected_diskseq,
        ["sfdisk", "--wipe", "always", devnode],
        input_text=script,
    )

    # Actualiser la table de partitions
    execute_device_command(
        devnode, expected_diskseq,
        ["partprobe", devnode]
    )

    execute_device_command(
        devnode, expected_diskseq,
        ["udevadm", "settle", "--timeout=15"],
        timeout=20
    )

    # Identifier la partition créée
    out = execute_device_command(
        devnode, expected_diskseq,
        ["lsblk", "-ln", "-o", "PATH,TYPE", devnode]
    )

    partitions = [
        parts[0]
        for line in out.splitlines()
        if len(parts := line.split()) == 2 and parts[1] == "part"
    ]

    if len(partitions) != 1:
        raise RuntimeError(
            f"Single partition not detected on {devnode}"
        )

    partition = partitions[0]

    # Formater la partition
    cmd = [cfg["tool"], *cfg["args"]]

    if label:
        cmd.extend([cfg["label_arg"], label])

    cmd.append(partition)

    execute_device_command(
        devnode, expected_diskseq,
        cmd,
        timeout=120
    )

    # Synchroniser les écritures
    execute_device_command(
        devnode, expected_diskseq,
        ["sync"],
        timeout=60
    )

    # Vérifier le système de fichiers
    detected = execute_device_command(
        devnode, expected_diskseq,
        ["blkid", "-p", "-o", "value", "-s", "TYPE", partition]
    ).strip().lower()

    if detected != cfg["blkid"]:
        raise RuntimeError(
            f"{fs_type.upper()} not detected after reformatting "
            f"(blkid : {detected or 'unknown'}, expected : {cfg['blkid']})"
        )

# ============================================================
# 6. CALCULS ET STATISTIQUES
# ============================================================

def compute_mbps(byte_count, duration):
    """Calcule le débit en MB/s à partir des octets transférés et de la durée."""
    if duration <= 0:
        return 0.0
    return byte_count / duration / 1e6


def compute_inter_cycle_stats(cycle_measurements):
    """Calcule la dispersion des débits moyens entre les cycles."""

    cycles = []

    for measurements in cycle_measurements:
        total_bytes = sum(size for size, duration in measurements)
        total_time = sum(duration for size, duration in measurements)

        if total_time > 0:
            cycles.append((
                compute_mbps(total_bytes, total_time),
                total_time,
            ))

    n = len(cycles)

    if n < 2:
        return {"n": n, "std": None, "cv": None}

    total_time = sum(duration for speed, duration in cycles)

    global_mean = sum(
        speed * duration for speed, duration in cycles
    ) / total_time

    variance = sum(
        duration * (speed - global_mean) ** 2
        for speed, duration in cycles
    ) / total_time

    std = float(np.sqrt(variance))

    return {
        "n": n,
        "std": std,
        "cv": std / global_mean * 100 if global_mean > 0 else None,
    }


def compute_intra_cycle_stats(cycle_measurements):
    """Calcule la dispersion des débits des blocs au sein des cycles."""

    total_bytes = 0
    total_time = 0.0
    weighted_variance = 0.0
    n = 0

    for measurements in cycle_measurements:
        cycle_bytes = sum(size for size, duration in measurements)
        cycle_time = sum(duration for size, duration in measurements)

        if cycle_time <= 0:
            continue

        cycle_mean = compute_mbps(cycle_bytes, cycle_time)

        for size, duration in measurements:
            if duration <= 0:
                continue

            speed = compute_mbps(size, duration)

            weighted_variance += duration * (speed - cycle_mean) ** 2
            n += 1

        total_bytes += cycle_bytes
        total_time += cycle_time

    if n < 2 or total_time <= 0:
        return {"n": n, "std": None, "cv": None}

    global_mean = compute_mbps(total_bytes, total_time)
    std = float(np.sqrt(weighted_variance / total_time))

    return {
        "n": n,
        "std": std,
        "cv": std / global_mean * 100 if global_mean > 0 else None,
    }

# ============================================================
# 7. GESTION DES RÉSULTATS ET DES ERREURS
# ============================================================

def create_test_result(devnode):
    """Initialise le dictionnaire des mesures et du verdict du test."""
    return {
        "devnode": devnode,
        "model": "N/A",
        "serial": "N/A",
        "csv": "",
        "docx": "",
        "report": "",
        "avg_write": 0.0,
        "avg_read": 0.0,
        "write_inter_stats": compute_inter_cycle_stats([]),
        "read_inter_stats": compute_inter_cycle_stats([]),
        "write_intra_stats": compute_intra_cycle_stats([]),
        "read_intra_stats": compute_intra_cycle_stats([]),
        "errors": 0,
        "error": "",
        "error_code": "",
        "error_detail": "",
        "result": "FAIL",
        "reformat": "",
        "completed": False,
    }


def classify_error(exc, operation=""):
    """Convertit une exception en erreur structurée pour le rapport."""

    if isinstance(exc, OSError):
        error_map = {
            errno.ENOENT: ("FILE_NOT_FOUND", "File or device not found"),
            errno.EIO: ("IO_ERROR", "Input/output error"),
            errno.ENODEV: ("DEVICE_UNAVAILABLE", "Device unavailable"),
            errno.ENXIO: ("DEVICE_UNAVAILABLE", "Device or address unavailable"),
            errno.ETIMEDOUT: ("IO_TIMEOUT", "Input/output timeout"),
            errno.EACCES: ("PERMISSION_DENIED", "Permission denied"),
            errno.EBUSY: ("DEVICE_BUSY", "Device busy"),
            errno.ENOSPC: ("NO_SPACE", "Insufficient disk space"),
            errno.EINVAL: ("INVALID_ARGUMENT", "Invalid argument"),
        }
        code, message = error_map.get(exc.errno, ("OS_ERROR", "System error"))

    elif isinstance(exc, ValueError):
        code, message = "INVALID_VALUE", "Invalid value"

    elif isinstance(exc, subprocess.TimeoutExpired):
        code, message = "COMMAND_TIMEOUT", "Command execution timeout"

    elif isinstance(exc, RuntimeError):
        code, message = "RUNTIME_ERROR", str(exc)

    else:
        code, message = "UNEXPECTED_ERROR", "Unexpected error"

    if operation:
        message = f"{operation} : {message}"

    return {
        "code": code,
        "message": message,
        "detail": f"{type(exc).__name__}: {exc}",
    }


def set_result_error(result, exc, operation=""):
    """Enregistre une exception et marque le test comme échoué."""
    error = classify_error(exc, operation)

    result["error"] = error["message"]
    result["error_code"] = error["code"]
    result["error_detail"] = error["detail"]
    result["result"] = "FAIL"
    result["completed"] = False

# ============================================================
# 8. GÉNÉRATION DES RAPPORTS
# ============================================================

def sanitize_filename(text):
    """Génère un nom de fichier avec des caractères ASCII."""
    return "".join(c if c.isascii() and (c.isalnum() or c in "-_") else "_"
                   for c in str(text))


def create_result_paths(serial):
    """Construit les chemins des fichiers de résultats du test USB."""
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    tag = f"{stamp}_{sanitize_filename(serial)}"

    return {
        "csv": str(RESULTS_DIR / f"USB_test_{tag}.csv"),
        "docx": str(RESULTS_DIR / f"USB_report_{tag}.docx"),
    }


def format_device_info(info):
    """Met en forme les informations du périphérique pour le rapport."""
    labels = (
        ("Date", "Date"),
        ("Model", "Model"),
        ("Manufacturer", "Manufacturer"),
        ("Serial", "Serial"),
        ("VID", "VID"),
        ("PID", "PID"),
        ("Size_GB", "Size_GB"),
        ("USBVersion", "USB Version"),
    )

    lines = ["=== DEVICE INFORMATION ==="]

    for key, label in labels:
        lines.append(f"{label:<14}: {info.get(key, 'N/A')}")

    return "\n".join(lines)


def get_failure_reasons(result):
    """Retourne les motifs de non-conformité du test."""
    reasons = []

    if result.get("error"):
        reasons.extend(result["error"].splitlines())

    if result["errors"]:
        reasons.append(f"{result['errors']} integrity error(s)")

    if 0 < result["avg_write"] <= MIN_WRITE_MBPS:
        reasons.append("Write speed below minimum threshold")

    if 0 < result["avg_read"] <= MIN_READ_MBPS:
        reasons.append("Read speed below minimum threshold")
    elif result["avg_read"] >= MAX_READ_MBPS:
        reasons.append("Read speed above maximum threshold")

    if not reasons:
        if result.get("reformat", "").startswith("FAILURE"):
            reasons.append("Reformatting failed (see REFORMATTING section)")
        else:
            reasons.append("Test incomplete or interrupted")

    return reasons


def build_report(info, result):
    """Construit le rapport textuel du test USB."""
    lines = [
        format_device_info(info),
        "",
        "=== PERFORMANCE ===",
        "",
    ]

    # Performances
    for label, avg, intra, inter in (
        (
            "Write",
            result["avg_write"],
            result["write_intra_stats"],
            result["write_inter_stats"],
        ),
        (
            "Read",
            result["avg_read"],
            result["read_intra_stats"],
            result["read_inter_stats"],
        ),
    ):
        lines.append(label)

        if intra["n"] > 0:
            lines.append(
                f"  {'Global average':<22}: {avg:8.2f} MB/s"
            )
        else:
            lines.append(
                f"  {'Global average':<22}: N/A"
            )

        if intra["std"] is not None:
            lines.append(
                f"  {'Intra-cycle STD':<22}: {intra['std']:8.2f} MB/s"
            )
            lines.append(
                f"  {'Intra-cycle CV':<22}: {intra['cv']:8.2f} %"
            )
        else:
            lines.append(
                f"  {'Intra-cycle STD / CV':<22}: N/A"
            )

        if inter["std"] is not None:
            lines.append(
                f"  {'Inter-cycle STD':<22}: {inter['std']:8.2f} MB/s"
            )
            lines.append(
                f"  {'Inter-cycle CV':<22}: {inter['cv']:8.2f} %"
            )
        else:
            lines.append(
                f"  {'Inter-cycle STD / CV':<22}: N/A"
            )

        lines.append("")

    # Intégrité
    lines.extend([
        "=== INTEGRITY ===",
        f"{'Integrity errors ':<20}: "
        f"{result['errors'] if result['read_intra_stats']['n'] > 0 else 'N/A'}",
        "",
    ])

    # Reformatage
    lines.extend([
        "=== REFORMATTING ===",
        result.get("reformat", "N/A"),
        "",
    ])

    # Verdict final
    lines.extend([
        "================= FINAL RESULT =================",
        "",
    ])

    if result["result"] == "PASS":
        lines.append("PASS")
    else:
        lines.append("FAIL")
        lines.extend(f"- {reason}" for reason in get_failure_reasons(result))

    lines.extend([
        "",
        "================================================",
    ])

    return "\n".join(lines)


def write_docx_report(
    template_path,
    output_path,
    report_text,
    figure,
    *,
    of_number=None,
    sap=None,
    device=None,
    serial=None,
    file_size=None,
    block_size=None,
    test_duration=None,
    final_fs=None,
    station="USB Tester",
    soft_version="TBD",
):
    """Génère le rapport DOCX avec le texte et le graphique du test."""

    template_path = Path(template_path)
    output_path = Path(output_path)

    if not template_path.is_file():
        raise FileNotFoundError(
            f"DOCX template not found  : {template_path}"
        )

    doc = Document(str(template_path))

    values = {
        "{{OF}}": of_number or "Not assigned",
        "{{SAP}}": sap or "Not assigned",
        "{{DEV_NAME}}": device or "N/A",
        "{{SERIAL}}": serial or "N/A",
        "{{FILE_SIZE}}": file_size or "N/A",
        "{{BLOCK_SIZE}}": block_size or "N/A",
        "{{TEST_DURATION}}": test_duration or "N/A",
        "{{FINAL_FS}}": final_fs or "N/A",
        "{{STATION}}": station,
        "{{SOFT_VERSION}}": soft_version,
    }

    found_text = False
    found_graph = False

    def replace_placeholder(paragraph, key, value):
        """Remplace un placeholder en conservant les runs non concernés."""

        # Cas courant : le placeholder est dans un seul run.
        for run in paragraph.runs:
            if key in run.text:
                run.text = run.text.replace(key, str(value))
                return True

        # Cas où Word a réparti le placeholder sur plusieurs runs.
        if key in paragraph.text:
            paragraph.text = paragraph.text.replace(key, str(value))
            return True

        return False

    def replace(paragraph):
        nonlocal found_text, found_graph

        if "{{RESULT_TEXT}}" in paragraph.text:
            found_text = True
            paragraph.clear()

            for index, line in enumerate(str(report_text).splitlines()):
                if index:
                    paragraph.add_run().add_break()

                paragraph.add_run(line)

            return

        if "{{RESULT_GRAPH}}" in paragraph.text:
            found_graph = True
            paragraph.clear()

            image = BytesIO()

            figure.savefig(
                image,
                format="png",
                dpi=180,
                bbox_inches="tight",
            )

            image.seek(0)

            paragraph.add_run().add_picture(
                image,
                width=Inches(5.8),
            )

            return

        for key, value in values.items():
            replace_placeholder(paragraph, key, value)

    def process_table(table):
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace(paragraph)

                for nested_table in cell.tables:
                    process_table(nested_table)

    for paragraph in doc.paragraphs:
        replace(paragraph)

    for table in doc.tables:
        process_table(table)

    if not found_text:
        raise RuntimeError(
            "Placeholder {{RESULT_TEXT}} not found in the template"
        )

    if not found_graph:
        raise RuntimeError(
            "Placeholder {{RESULT_GRAPH}} not found in the template"
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)

    # Enregistrer d'abord dans un fichier temporaire.
    temp_path = output_path.with_name(
        output_path.stem + ".tmp.docx"
    )

    try:
        doc.save(str(temp_path))
        temp_path.replace(output_path)

    finally:
        temp_path.unlink(missing_ok=True)

    # Conversion DOCX -> PDF avec LibreOffice.
    subprocess.run(
        [
            "libreoffice",
            "--headless",
            "--convert-to", "pdf",
            "--outdir", str(output_path.parent),
            str(output_path),
        ],
        check=True,
        timeout=60,
        capture_output=True,
        text=True,
    )

    return output_path

# ============================================================
# 8. GESTION DES COMPTEURS D'USURE DES TEST-LINKS
# ============================================================

def load_test_link_counters():
    """Load persistent mating cycle counters."""

    if not TEST_LINK_FILE.exists():
        return {}

    with open(TEST_LINK_FILE, "r", encoding="utf-8") as file:
        return json.load(file)


def save_test_link_counters(counters):
    """Save mating cycle counters atomically."""

    TEST_LINK_FILE.parent.mkdir(parents=True, exist_ok=True)

    temporary_path = None

    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=TEST_LINK_FILE.parent,
            delete=False,
        ) as file:
            temporary_path = file.name

            os.chmod(file.name, 0o644)

            json.dump(counters, file, indent=4)
            file.flush()
            os.fsync(file.fileno())

        os.replace(temporary_path, TEST_LINK_FILE)

    finally:
        if temporary_path and os.path.exists(temporary_path):
            os.unlink(temporary_path)


def increment_test_link_counter(test_link):
    """Increment and persist the mating cycle counter."""

    counters = load_test_link_counters()

    if test_link not in counters:
        counters[test_link] = {
            "cycles": 0,
            "max_cycles": MAX_MATING_CYCLES,
        }

    counters[test_link]["cycles"] += 1

    save_test_link_counters(counters)

    return counters[test_link]["cycles"]